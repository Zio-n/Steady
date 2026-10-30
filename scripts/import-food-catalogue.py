#!/usr/bin/env python3
"""Import source-attributed CoFID generic foods into a running local Supabase catalogue.

Requires openpyxl and requests. Source files are downloaded separately; do not commit them.
"""
import argparse
import base64
import hashlib
import hmac
import json
import os
import re
import subprocess
import sys
import uuid
from collections import defaultdict

import requests
from openpyxl import load_workbook

BASE = "http://127.0.0.1:54331/rest/v1"
SOURCE = {
    "id": "cofid_2021",
    "name": "CoFID 2021",
    "licence": "Open Government Licence v3.0",
    "source_url": "https://www.gov.uk/government/publications/composition-of-foods-integrated-dataset-cofid",
    "attribution": "Contains public sector information licensed under the Open Government Licence v3.0.",
}


def slug(value):
    value = value.lower().replace("µ", "micro").replace("%", "percent")
    value = re.sub(r"[^a-z0-9]+", "_", value).strip("_")
    value = (value or "nutrient")[:64]
    return value if value[0].isalpha() else "n_" + value[:62]


def numeric(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        value = value.strip().replace(",", "")
        if not value or value.lower() in {"tr", "trace", "n", "na", "n/a", "-", "—"}:
            return None
        try:
            return float(value)
        except ValueError:
            return None
    return None


COFID_COMMON = {
    "WATER": ("water_g", "Water", "g"),
    "TOTNIT": ("total_nitrogen_g", "Total nitrogen", "g"),
    "PROT": ("protein_g", "Protein", "g"),
    "FAT": ("fat_g", "Fat", "g"),
    "CHO": ("carbohydrate_g", "Carbohydrate", "g"),
    "KCALS": ("energy_kcal", "Energy", "kcal"),
    "KJ": ("energy_kj", "Energy", "kJ"),
    "STAR": ("starch_g", "Starch", "g"),
    "OLIGO": ("oligosaccharides_g", "Oligosaccharides", "g"),
    "TOTSUG": ("sugars_g", "Sugars", "g"),
    "AOACFIB": ("fibre_g", "Dietary fibre", "g"),
    "SATFOD": ("saturated_fat_g", "Saturated fat", "g"),
    "FODTRANS": ("trans_fat_g", "Trans fat", "g"),
    "CHOL": ("cholesterol_mg", "Cholesterol", "mg"),
    "NA": ("sodium_mg", "Sodium", "mg"),
    "K": ("potassium_mg", "Potassium", "mg"),
    "CA": ("calcium_mg", "Calcium", "mg"),
    "MG": ("magnesium_mg", "Magnesium", "mg"),
    "P": ("phosphorus_mg", "Phosphorus", "mg"),
    "FE": ("iron_mg", "Iron", "mg"),
    "CU": ("copper_mg", "Copper", "mg"),
    "ZN": ("zinc_mg", "Zinc", "mg"),
    "VITC": ("vitamin_c_mg", "Vitamin C", "mg"),
}
COFID_ALCOHOL_GROUPS = {"Q", "QA", "QC", "QE", "QF", "QG", "QI", "QK"}


def cofid_value(value):
    if isinstance(value, str) and value.strip().lower() in {"tr", "trace"}:
        return 0.0, True
    return numeric(value), False


def parse_cofid(path):
    book = load_workbook(path, read_only=True, data_only=True)
    foods = {}
    for sheet in book.worksheets:
        if not re.match(r"1\.\d+", sheet.title) or sheet.title.endswith("Notes"):
            continue
        # 1.2 contains food conversion/factor fields, while the FA tables are
        # expressed per 100 g fatty acid rather than per 100 g food.
        if sheet.title.startswith("1.2 ") or re.search(r"per 100g?FA", sheet.title, re.IGNORECASE):
            continue
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, ())
        codes = next(rows, ())
        next(rows, None)  # human-readable nutrient headings
        for row in rows:
            if not row or not row[0] or not row[1]:
                continue
            source_key = str(row[0]).strip()
            name = str(row[1]).strip()
            group = str(row[3] or "").strip() if len(row) > 3 else ""
            item = foods.setdefault(source_key, {
                "id": str(uuid.uuid5(uuid.NAMESPACE_URL, "cofid:" + source_key)),
                "source_id": SOURCE["id"], "source_record_id": source_key,
                "data_kind": "generic", "name": name, "brand": None,
                "country_code": "GB", "locale": "en-GB",
                "basis_amount": 100, "basis_unit": "ml" if group in COFID_ALCOHOL_GROUPS else "g",
                "quality_status": "incomplete", "source_metadata": {"cofid_group": group or None,
                    "trace_values": []},
                "nutrients": {},
            })
            for idx in range(7, min(len(row), len(headers))):
                value, is_trace = cofid_value(row[idx])
                if value is None:
                    continue
                heading = str(headers[idx] or "").strip()
                if not heading:
                    continue
                raw_code = str(codes[idx] or "").strip()
                if not raw_code:
                    continue
                # The main proximates sheet also contains fatty-acid ratios
                # per 100 g of fatty acid; those are not per 100 g of food.
                if re.search(r"/\s*100g\s*FA\b", heading, re.IGNORECASE) or "FAC" in raw_code.upper():
                    continue
                unit_match = re.search(r"\(([^()]*)\)\s*$", heading)
                unit = unit_match.group(1).strip().lower().replace("µ", "µ") if unit_match else "g"
                common = COFID_COMMON.get(raw_code.upper())
                if common:
                    code, name, unit = common
                else:
                    code = "cofid_" + slug(raw_code)[:58]
                    name = heading
                # Common mineral units are defined in mg; source columns may be µg.
                if code in {"sodium_mg", "potassium_mg", "calcium_mg", "magnesium_mg", "phosphorus_mg", "iron_mg", "copper_mg", "zinc_mg"} and unit in {"µg", "ug", "mcg"}:
                    value /= 1000
                    unit = "mg"
                if is_trace:
                    item["source_metadata"]["trace_values"].append(raw_code)
                item["nutrients"][code] = {"value": value, "unit": unit, "name": heading}
    for item in foods.values():
        item["source_metadata"]["trace_values"] = sorted(set(item["source_metadata"]["trace_values"]))
        item["quality_status"] = "usable" if "energy_kcal" in item["nutrients"] else "incomplete"
    book.close()
    return list(foods.values())


def chunks(items, size=200):
    for start in range(0, len(items), size):
        yield items[start:start + size]


def local_service_jwt(secret):
    def encode(part):
        return base64.urlsafe_b64encode(json.dumps(part, separators=(",", ":")).encode()).rstrip(b"=").decode()
    header = encode({"alg": "HS256", "typ": "JWT"})
    payload = encode({"role": "service_role", "iss": "supabase", "iat": 1793350000, "exp": 1983812996})
    signature = base64.urlsafe_b64encode(hmac.new(secret.encode(), f"{header}.{payload}".encode(), hashlib.sha256).digest()).rstrip(b"=").decode()
    return f"{header}.{payload}.{signature}"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--cofid", required=True)
    args = parser.parse_args()
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if not key and os.environ.get("SUPABASE_JWT_SECRET"):
        key = local_service_jwt(os.environ["SUPABASE_JWT_SECRET"])
    if not key:
        status = subprocess.run(["npx", "supabase", "status"], capture_output=True, text=True, check=True)
        key = json.loads(status.stdout)["SERVICE_ROLE_KEY"]
    session = requests.Session()
    session.headers.update({"apikey": key, "Authorization": "Bearer " + key, "Content-Type": "application/json"})

    response = session.post(BASE + "/catalog_sources?on_conflict=id", json=SOURCE,
                            headers={"Prefer": "resolution=merge-duplicates,return=minimal"}, timeout=30)
    response.raise_for_status()
    foods = parse_cofid(args.cofid)
    definitions = {}
    for food in foods:
        for code, nutrient in food["nutrients"].items():
            definitions.setdefault(code, {"nutrient_code": code, "name": nutrient["name"],
                                           "standard_unit": nutrient["unit"], "sort_order": 500})
    for batch in chunks(list(definitions.values())):
        r = session.post(BASE + "/catalog_nutrient_definitions?on_conflict=nutrient_code", json=batch,
                         headers={"Prefer": "resolution=merge-duplicates,return=minimal"}, timeout=30)
        if not r.ok:
            raise RuntimeError(f"Nutrient definition import failed: {r.text}")
        r.raise_for_status()

    count = 0
    for batch in chunks(foods, 100):
        rows = [{k: v for k, v in food.items() if k != "nutrients"} for food in batch]
        r = session.post(BASE + "/catalog_foods?on_conflict=source_id,source_record_id", json=rows,
                         headers={"Prefer": "resolution=merge-duplicates,return=representation"}, timeout=60)
        r.raise_for_status()
        returned = {row["source_record_id"]: row["id"] for row in r.json()}
        nutrient_rows = []
        for food in batch:
            food_id = returned.get(food["source_record_id"], food["id"])
            for code, nutrient in food["nutrients"].items():
                nutrient_rows.append({"food_id": food_id, "nutrient_code": code,
                    "value": nutrient["value"], "source_unit": nutrient["unit"],
                    "basis_amount": 100, "basis_unit": food["basis_unit"], "source_nutrient_name": nutrient["name"]})
        for nb in chunks(nutrient_rows):
            r = session.post(BASE + "/catalog_food_nutrients?on_conflict=food_id,nutrient_code", json=nb,
                             headers={"Prefer": "resolution=merge-duplicates,return=minimal"}, timeout=60)
            r.raise_for_status()
        count += len(batch)
        if count % 500 == 0 or count == len(foods):
            print(f"Imported {count}/{len(foods)} CoFID foods", flush=True)
    print(f"CoFID import complete: {len(foods)} foods, {len(definitions)} nutrient definitions")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"Import failed: {error}", file=sys.stderr)
        sys.exit(1)
